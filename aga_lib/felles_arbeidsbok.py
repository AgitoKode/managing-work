"""Bygger ÉN samlet arbeidsbok som dekker alle steg (STEG 1-9) i oppdraget.

Gjenbruker aga_analyse.beregn_grunnlag() og aga_lib.leveranser_steg4_9 for selve
tallgrunnlaget - denne modulen inneholder kun sammenstilling og
Excel-formatering, slik at tallene alltid er identiske med de øvrige
leveransene (AGA_Rapport_<år>.xlsx og de seks STEG4-9-filene).
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from aga_lib import report
from aga_lib.leveranser_steg4_9 import bygg_leveranse_data

FONT = "Arial"


def _unik_sti(sti: Path) -> Path:
    if not sti.exists():
        return sti
    tidsstempel = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    return sti.with_name(f"{sti.stem}_{tidsstempel}{sti.suffix}")


def _tall_tekst(v) -> str:
    if isinstance(v, float) and v == v and v.is_integer():
        return str(int(v))
    if v is None or (isinstance(v, float) and v != v):
        return ""
    return str(v)


# ---------------------------------------------------------------------------
# STEG 3 - kommune/AGA-sone-dokumentasjon (samme kuraterte innhold som i den
# frittstående AGA-soner-satser-2026-dokumentasjon.xlsx / STEG1-3-filen)
# ---------------------------------------------------------------------------
_SKD_TITTEL = "Satser for arbeidsgiveravgift - soneinndeling (Kommunekatalog 2026)"
_SKD_URL = "https://www.skatteetaten.no/satser/arbeidsgiveravgift-soneinndeling/#kommunekatalog-2026"
_SKD_UTGIVER = "Skatteetaten"
_SKD_DATO_PUB = "Ikke oppgitt i uttrekket - bør bekreftes mot sidefot på kildesiden"

_SATS = {"1": "14,1 %", "2": "10,6 %", "5": "0 %"}

_STEG3_RADER = [
    dict(kommune="Oslo", knr="0301", lok="0469, 0319, 0668, 0440 (flere arbeidssteder i Oslo)",
         sone="1", sats=_SATS["1"],
         merknad="Hele Oslo kommune er sone 1. Ingen kjent soneinndeling internt i kommunen.",
         status="KILDEVERIFISERT"),
    dict(kommune="Drammen", knr="3301", lok="3041", sone="1", sats=_SATS["1"],
         merknad="Drammen (sammenslått 2020) står som samlet sone-1-kommune i Kommunekatalog 2026.",
         status="KILDEVERIFISERT"),
    dict(kommune="Hole", knr="3310", lok="3530 (Røyse)", sone="1", sats=_SATS["1"],
         merknad="Hele Hole kommune er sone 1.", status="KILDEVERIFISERT"),
    dict(kommune="Ullensaker", knr="3209",
         lok="Postnummer ikke oppgitt; kommune identifisert via bedriftsnavn «Ullensaker kommune»",
         sone="1", sats=_SATS["1"],
         merknad="Hele Ullensaker kommune er sone 1 - ingen intern deling.", status="KILDEVERIFISERT"),
    dict(kommune="Vindafjord", knr="1160", lok="5580 (Ølen)", sone="1", sats=_SATS["1"],
         merknad="Hele Vindafjord kommune er sone 1.", status="KILDEVERIFISERT"),
    dict(kommune="Sauda", knr="1135", lok="4201", sone="2", sats=_SATS["2"],
         merknad="Hele Sauda kommune er sone 2.", status="KILDEVERIFISERT"),
    dict(kommune="Sveio", knr="4612", lok="5550", sone="1", sats=_SATS["1"],
         merknad="Hele Sveio kommune er sone 1.", status="KILDEVERIFISERT"),
    dict(kommune="Sunnfjord", knr="4647", lok="6800 (Førde - tidligere Førde kommune)", sone="1a",
         sats="10,6 % inntil fribeløpet (kr 850 000/foretak i 2026), deretter 14,1 %",
         merknad="Sunnfjord er DELT: «Gamle Førde» = sone 1a, «Gamle Gaular, Jølster og Naustdal» = sone 2. "
                 "Postnummer 6800 Førde => sone 1a. Se egen rad under for kommunens øvrige areal.",
         status="KILDEVERIFISERT"),
    dict(kommune="Sunnfjord (øvrig areal - ikke brukt i dette datagrunnlaget)", knr="4647",
         lok="Gamle Gaular, Jølster og Naustdal", sone="2", sats=_SATS["2"],
         merknad="Dokumentert for fullstendighet - ingen arbeidssteder i datagrunnlaget ligger her.",
         status="KILDEVERIFISERT"),
    dict(kommune="Kvænangen", knr="5546", lok="9161 (Burfjord)", sone="5", sats=_SATS["5"],
         merknad="Hele Kvænangen kommune er sone 5 (0 %, tiltakssonen).", status="KILDEVERIFISERT"),
    dict(kommune="Skjervøy", knr="5542",
         lok="Postnummer ikke oppgitt; kommune identifisert via bedriftsnavn «Skjervøy kommune hjemmetjenesten»",
         sone="5", sats=_SATS["5"], merknad="Hele Skjervøy kommune er sone 5 (0 %).",
         status="KILDEVERIFISERT"),
    dict(kommune="Nordreisa", knr="5544",
         lok="«Sonjatun Sykestue» (Storslett) - MEN RecMan sitt prosjektnavn sier «...Kvænangen kommune»",
         sone="5", sats=_SATS["5"],
         merknad="MOTSTRIDENDE SIGNALER: institusjonen ligger fysisk i Nordreisa, men prosjektnavnet i "
                 "RecMan sier Kvænangen. Begge er sone 5 (0 %), så AGA-satsen er upåvirket, men "
                 "kommunetilhørigheten bør avklares med lønn/RecMan. Se STEG6/Avviksrapport.",
         status="MÅ KVALITETSSIKRES – motstridende opplysninger om kommune"),
]


def _bygg_steg1(hovedark, arknavn: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    df_raw = hovedark.df
    mapping = hovedark.kolonnemapping

    kolonnerader = []
    for kol in df_raw.columns:
        serie = df_raw[kol]
        ikke_null = serie.dropna()
        observert_type = "(ingen data)"
        if len(ikke_null):
            typer = {type(v).__name__ for v in ikke_null.head(200)}
            observert_type = ", ".join(sorted(typer))
        kolonnerader.append({
            "Original kolonne": kol,
            "Standardfelt": mapping.original_til_standard.get(kol, "(ikke koblet)"),
            "Observert datatype": observert_type,
            "Antall rader": len(serie),
            "Manglende verdier": int(serie.isna().sum()),
            "Unike verdier": int(serie.nunique(dropna=True)),
        })
    kolonner_df = pd.DataFrame(kolonnerader)

    dato_kolonner = ["Dato", "Første dag", "Siste dag", "Dato registrert", "Jobb start", "Jobb slutt", "Jobb opprettet", "Fakturadato"]
    from aga_lib.dates import til_dato
    ugyldige_total = 0
    dato_funn = []
    for kol in dato_kolonner:
        if kol not in df_raw.columns:
            continue
        ikke_null = df_raw[kol].dropna()
        ugyldige = [v for v in ikke_null if til_dato(v) is None]
        ugyldige_total += len(ugyldige)
        dato_funn.append(f"{kol}: {len(ugyldige)} ugyldige av {len(ikke_null)} utfylte")

    datakvalitet_df = pd.DataFrame([
        {"Kategori": "Antall rader importert", "Funn": str(len(df_raw)), "Alvorlighet": "Info",
         "Kommentar": f"Ark {arknavn!r} valgt automatisk som hoveddataark."},
        {"Kategori": "Fullstendige radduplikater", "Funn": str(int(df_raw.duplicated().sum())), "Alvorlighet": "Info",
         "Kommentar": "Rader identiske i ALLE kolonner. Tilleggslinjer per vakt er IKKE duplikater."},
        {"Kategori": "Ugyldige/utolkbare datoer", "Funn": str(ugyldige_total), "Alvorlighet": "Info",
         "Kommentar": "; ".join(dato_funn)},
        {"Kategori": "Kolonner lagret som tekst men ser numeriske ut",
         "Funn": ", ".join(hovedark.kolonner_lagret_som_tekst_men_ser_numerisk_ut) or "Ingen", "Alvorlighet": "Info",
         "Kommentar": "Postnummer er BEVISST tekst for å bevare ledende null - skal ikke konverteres."},
        {"Kategori": "Manglende Postnummer", "Funn": str(int(df_raw["Postnummer"].isna().sum())) if "Postnummer" in df_raw.columns else "0",
         "Alvorlighet": "Middels", "Kommentar": "Se STEG2/STEG6 for hvordan kommune ble identifisert uten postnummer."},
    ])
    return kolonner_df, datakvalitet_df


def _bygg_steg2(df: pd.DataFrame) -> pd.DataFrame:
    rader = []
    for prosjektnr, g in df.groupby("prosjektnr", dropna=False):
        prosjekt_navn = sorted({str(v) for v in g["prosjekt"].dropna().unique()})
        bedrift = sorted({str(v) for v in g["bedrift"].dropna().unique()})
        org = sorted({_tall_tekst(v) for v in g["org_nr"].dropna().unique()} - {""})
        postnr = sorted({str(v) for v in g["postnummer"].dropna().unique()})
        rader.append({
            "Prosjektnr": prosjektnr,
            "Prosjekt": "; ".join(prosjekt_navn),
            "Kommune": (g["kommune"].dropna().unique().tolist() or [""])[0],
            "AGA-sone": (g["sone"].dropna().unique().tolist() or [""])[0],
            "Kunde": "; ".join(bedrift),
            "Bedrift": "; ".join(bedrift),
            "Org.nr": "; ".join(org),
            "Postnummer": "; ".join(postnr),
            "Arbeidstimer": round(float(g["arbeidstimer"].sum()), 2),
            "Total lønn (kr)": round(float(pd.to_numeric(g["total_lonn"], errors="coerce").sum()), 2),
            "Antall rader": len(g),
        })
    return pd.DataFrame(rader)


def _bygg_steg7(df: pd.DataFrame) -> list[tuple[str, pd.DataFrame]]:
    df = df.copy()
    df["total_lonn"] = pd.to_numeric(df.get("total_lonn"), errors="coerce")
    df["lonn_inkl_sosial_kost"] = pd.to_numeric(df.get("lonn_inkl_sosial_kost"), errors="coerce")

    def agg(gruppefelt: str, kolonnenavn: str) -> pd.DataFrame:
        g = df.groupby(gruppefelt, dropna=False)
        resultat = g.agg(
            Antall_ansatte=("ansattnr", "nunique"),
            Timer=("arbeidstimer", "sum"),
            Total_lonn=("total_lonn", "sum"),
            Total_lonn_inkl_sosial_kost=("lonn_inkl_sosial_kost", "sum"),
        ).reset_index().rename(columns={
            gruppefelt: kolonnenavn, "Antall_ansatte": "Antall ansatte",
            "Total_lonn": "Total lønn", "Total_lonn_inkl_sosial_kost": "Total lønn inkl. sosial kost",
        })
        return resultat

    return [
        ("Per kommune", agg("kommune", "Kommune")),
        ("Per prosjekt", agg("prosjektnr", "Prosjektnr")),
        ("Per kunde/bedrift", agg("bedrift", "Kunde/Bedrift")),
        ("Per AGA-sone", agg("sone", "AGA-sone")),
        ("Per måned", agg("maaned", "Måned")),
        ("Per år", agg("aar", "År")),
    ]


def bygg_felles_arbeidsbok(prosjektmappe: Path, grunnlag: dict | None = None) -> Path:
    """Samler STEG 1-9 i én arbeidsbok: .\\output\\AGA-analyse-<år>-alle-steg.xlsx.

    Tar imot et allerede beregnet `grunnlag` (fra aga_analyse.beregn_grunnlag())
    når denne kalles som del av den samlede kjøringen i kjoer_analyse(), slik at
    kildefilen kun leses og analyseres ÉN gang per prosess. Kalles denne
    frittstående (grunnlag=None), beregnes grunnlaget her på vanlig måte."""
    import aga_analyse  # lokal import for å unngå sirkulær import ved modulnivå

    if grunnlag is None:
        grunnlag = aga_analyse.beregn_grunnlag(prosjektmappe)
    konfig = grunnlag["konfig"]
    logger = grunnlag["logger"]
    df = grunnlag["df"]
    analyseaar = grunnlag["analyseaar"]
    lesekontekst = grunnlag["lesekontekst"]
    hovedark = lesekontekst["alle_ark"][lesekontekst["arknavn"]]

    leveransedata = bygg_leveranse_data(
        df=df,
        prosjektregister=grunnlag["prosjektregister"],
        aga_per_prosjektnoekkel=grunnlag["aga_per_prosjektnoekkel"],
        kilde_id_geografi=grunnlag["kilde_id_geografi"],
        kilde_tilgjengelig=grunnlag["kilde_tilgjengelig"],
        kk_url=_SKD_URL,
        kk_tittel=_SKD_TITTEL,
        kk_utgiver=_SKD_UTGIVER,
        kontrolltidspunkt=grunnlag["kontrolltidspunkt"],
        analyseaar=analyseaar,
        logger=logger,
    )

    wb = report.lag_arbeidsbok()

    # --- STEG 1 ---
    kolonner_df, datakvalitet1_df = _bygg_steg1(hovedark, lesekontekst["arknavn"])
    report.skriv_tabellark(wb, "STEG1 - Kolonner og datatyper", kolonner_df,
                            undertittel="STEG 1 - Kolonner, datatyper og manglende verdier")
    report.skriv_tabellark(wb, "STEG1 - Datakvalitet", datakvalitet1_df,
                            undertittel="STEG 1 - Dokumenterte datakvalitetsfunn")

    # --- STEG 2 ---
    steg2_df = _bygg_steg2(df)
    report.skriv_tabellark(wb, "STEG2 - Prosjektregister", steg2_df,
                            undertittel="STEG 2 - Ett unikt prosjekt per rad, med kommune/AGA-sone og arbeidstimer/lønn")

    # --- STEG 3 ---
    steg3_rows = []
    for r in _STEG3_RADER:
        steg3_rows.append({
            "År": analyseaar, "Kommune": r["kommune"], "Kommunenummer": r["knr"],
            "Postnummer/lokasjon": r["lok"], "AGA-sone": r["sone"], "AGA-sats": r["sats"],
            "Kildetittel": _SKD_TITTEL, "Kilde-URL": _SKD_URL, "Offentlig utgiver": _SKD_UTGIVER,
            "Publisert/oppdatert dato": _SKD_DATO_PUB, "Regelverksmerknad": r["merknad"],
            "Kontrollstatus": r["status"],
        })
    report.skriv_tabellark(
        wb, "STEG3 - Kildedokumentasjon", pd.DataFrame(steg3_rows), status_kolonner=["Kontrollstatus"],
        undertittel="STEG 3 - AGA-soner og -satser for 2026, med kildedokumentasjon",
    )
    sats_df = pd.DataFrame([
        {"Sone": "I", "Ordinære næringer": "14,1 %", "Landbruk og fiske": "14,1 %"},
        {"Sone": "Ia*", "Ordinære næringer": "14,1 % (10,6 % inntil fribeløp)", "Landbruk og fiske": "10,6 %"},
        {"Sone": "II", "Ordinære næringer": "10,6 %", "Landbruk og fiske": "10,6 %"},
        {"Sone": "III", "Ordinære næringer": "6,4 %", "Landbruk og fiske": "6,4 %"},
        {"Sone": "IV", "Ordinære næringer": "5,1 %", "Landbruk og fiske": "5,1 %"},
        {"Sone": "IVa", "Ordinære næringer": "7,9 %", "Landbruk og fiske": "5,1 %"},
        {"Sone": "V", "Ordinære næringer": "0 %", "Landbruk og fiske": "0 %"},
    ])
    report.skriv_tabellark(wb, "STEG3 - Satstabell 2026", sats_df,
                            undertittel="* Sone Ia: 10,6 % inntil fribeløpet (kr 850 000/foretak i 2026) er brukt opp, deretter 14,1 %.")

    # --- STEG 4 (metodenotat - selve sone-oppslaget vises i STEG3/STEG5) ---
    steg4_linjer = [
        ("PRIMÆRKILDE", "Skatteetaten"),
        ("SEKUNDÆRKILDE", "Lovdata"),
        ("TERTIÆRKILDE", "Regjeringen.no"),
        ("IKKE BRUKT", "Blogger, forum, KI-genererte oversikter, private regnskapsnettsteder"),
        ("REGEL VED MANGLENDE DOKUMENTASJON", "Status settes til «KILDE IKKE FUNNET» - ingen gjetting er tillatt."),
        (
            "STATUS I DENNE KJØRINGEN",
            "skatteetaten.no/lovdata.no/regjeringen.no var ikke nåbare fra dette kjøremiljøet (se "
            "nettverksprobe i Kjøremetadata-arket i AGA_Rapport). Kommunekatalog 2026 og satstabellen "
            "ble i stedet mottatt direkte fra oppdragsgiver som utdrag av Skatteetatens egen side, "
            "med full kildehenvisning (se STEG3/STEG5).",
        ),
        ("RESULTAT", "Se STEG3 - Kildedokumentasjon og STEG5 - Kilderegister for AGA-sone per identifisert kommune."),
    ]
    report.skriv_lederoppsummering(wb, "STEG4 - Kildehierarki", steg4_linjer)

    # --- STEG 5 ---
    report.skriv_tabellark(
        wb, "STEG5 - Kilderegister", leveransedata["kilderegister"],
        status_kolonner=["Status"] if len(leveransedata["kilderegister"]) else None,
        undertittel="STEG 5 - Kilderegister: Kilde-ID | Kommune | AGA-sone | Kilde | URL | Utgiver | Kontrollert dato | Status",
    )

    # --- STEG 6 ---
    report.skriv_tabellark(
        wb, "STEG6 - Detaljgrunnlag", leveransedata["detaljrapport"],
        undertittel="STEG 6 - Arbeid -> Prosjektnr -> Kommune -> AGA-sone, én rad per registrering",
    )

    # --- STEG 7 ---
    ws7 = wb.create_sheet("STEG7 - Aggregering")
    ws7.cell(row=1, column=1, value="STEG 7 - Summering per kommune, prosjekt, kunde, AGA-sone, måned og år").font = Font(name=FONT, size=13, bold=True)
    rad = 3
    for tittel, delta_df in _bygg_steg7(df):
        ws7.cell(row=rad, column=1, value=tittel).font = Font(name=FONT, size=11, bold=True)
        rad += 1
        header_rad = rad
        for j, kol in enumerate(delta_df.columns, start=1):
            c = ws7.cell(row=header_rad, column=j, value=str(kol))
            c.font = Font(name=FONT, size=10, bold=True)
        for i, (_, r) in enumerate(delta_df.iterrows(), start=header_rad + 1):
            for j, kol in enumerate(delta_df.columns, start=1):
                verdi = r[kol]
                if isinstance(verdi, float) and pd.isna(verdi):
                    verdi = None
                ws7.cell(row=i, column=j, value=verdi).font = Font(name=FONT, size=10)
        rad = header_rad + len(delta_df) + 2
    for kol_idx in range(1, 6):
        ws7.column_dimensions[get_column_letter(kol_idx)].width = 22

    # --- STEG 8 ---
    report.skriv_tabellark(
        wb, "STEG8 - Kontrollrapport", leveransedata["avviksrapport"],
        undertittel="STEG 8 - Kontrollrapport/avvik: prosjekter uten kommune, kommuner uten AGA-sone, "
        "manglende postnummer, motstridende datagrunnlag, flere kommuner/soner per prosjekt",
    )

    # --- STEG 9 ---
    md = leveransedata["lederoppsummering_md"]
    seksjoner = [s.strip() for s in md.split("## ") if s.strip()][1:]  # dropp H1-tittelen
    steg9_linjer = []
    for seksjon in seksjoner:
        linjer = seksjon.split("\n", 1)
        tittel = linjer[0].strip()
        innhold = linjer[1].strip() if len(linjer) > 1 else ""
        steg9_linjer.append((tittel, innhold))
    report.skriv_lederoppsummering(wb, "STEG9 - Ledelsesrapport", steg9_linjer)

    # --- Metode og forbehold ---
    metode_linjer = [
        ("Om denne arbeidsboken", True, 13),
        ("", False, 10),
        (
            "Denne arbeidsboken samler alle ni steg i oppdraget i én fil. Samme underliggende "
            "beregning brukes også i AGA_Rapport_<år>.xlsx (15 ark) og de seks separate "
            "STEG4-9-filene i .\\output - tallene er identiske på tvers av alle tre leveranseformer.",
            False, 10,
        ),
        ("", False, 10),
        (
            "Viktig forbehold: Ingen AGA-sone i denne arbeidsboken er endelig juridisk bekreftet. "
            "Det er ikke avklart mot primærkilde om arbeidsstedets kommune (fremfor f.eks. registrert "
            "underenhet) er riktig AGA-basis for et bemanningsforetak som leier ut arbeidskraft. "
            "Se STEG4/STEG9 og Regelverksgrunnlag-arket i AGA_Rapport_<år>.xlsx for detaljer.",
            False, 10,
        ),
        (
            "Skatteetaten.no, lovdata.no og regjeringen.no var ikke nåbare fra dette kjøremiljøet ved "
            "kjøretidspunktet - kildene i STEG5 - Kilderegister er derfor ikke selv åpnet og lest i "
            "denne kjøringen. Kommunekatalog 2026 og satstabellen ble mottatt direkte fra "
            "oppdragsgiver som utdrag av Skatteetatens egne sider (se STEG3/STEG4).",
            False, 10,
        ),
    ]
    ws_metode = wb.create_sheet("Metode og forbehold")
    ws_metode.column_dimensions["A"].width = 112
    r = 1
    for tekst, fet, storrelse in metode_linjer:
        c = ws_metode.cell(row=r, column=1, value=tekst)
        c.font = Font(name=FONT, size=storrelse, bold=fet, color="1F4E78" if fet else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws_metode.row_dimensions[r].height = 48 if len(tekst) > 90 else (18 if tekst else 8)
        r += 1

    output_mappe = konfig.sti("output_mappe")
    output_sti = _unik_sti(output_mappe / f"AGA-analyse-{analyseaar}-alle-steg.xlsx")
    report.lagre(wb, output_sti)
    logger.info("Felles arbeidsbok for alle steg lagret: %s", output_sti)
    return output_sti
