"""Bygger AGA_Rapport_2026.xlsx med alle 15 påkrevde ark.

Fellesregler for alle datatabellark (seksjon 17):
- fet, hvit overskriftsrad med farget bakgrunn
- autofilter og frosset overskriftsrad
- kolonnebredde tilpasset innholdet
- ekte Excel-tabell (ListObject)
- dato som dd.mm.åååå, beløp som norsk tallformat med to desimaler
- fargekoding av kontrollstatus (grønn/gul/rød/grå)
- ingen sammenslåtte celler, ingen makroer, ikke passordbeskyttet
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Arial"
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
BODY_FONT = Font(name=FONT, size=10)
TITTEL_FONT = Font(name=FONT, size=13, bold=True)
UNDERTITTEL_FONT = Font(name=FONT, size=10, italic=True)

FARGE_GRONN = PatternFill("solid", fgColor="C6E0B4")
FARGE_GUL = PatternFill("solid", fgColor="FFE699")
FARGE_ROD = PatternFill("solid", fgColor="F8CBAD")
FARGE_GRAA = PatternFill("solid", fgColor="D9D9D9")

BELOPFORMAT = "#,##0.00"
PROSENTFORMAT = "0.0%"
HELTALLFORMAT = "#,##0"


def _statusfarge(verdi: str) -> PatternFill | None:
    if not verdi:
        return None
    v = str(verdi).upper()
    if "IKKE_FUNNET" in v or "MOTSTRIDENDE" in v or "KILDE_IKKE_FUNNET" in v:
        return FARGE_ROD
    if "IKKE_RELEVANT" in v:
        return FARGE_GRAA
    if v.endswith("_VERIFISERT") or v == "KILDEVERIFISERT":
        return FARGE_GRONN
    if "MÅ_KVALITETSSIKRES" in v or "SANNSYNLIG" in v or "GEOGRAFISK ANALYSEGRUNNLAG" in v:
        return FARGE_GUL
    return None


def _sikkert_arknavn(navn: str) -> str:
    return navn[:31]


def skriv_tabellark(
    wb: Workbook,
    arknavn: str,
    df: pd.DataFrame,
    kolonnetyper: dict[str, str] | None = None,
    status_kolonner: list[str] | None = None,
    undertittel: str | None = None,
) -> None:
    """Skriver en DataFrame som et formatert Excel-tabellark. kolonnetyper kan
    angi 'dato', 'beloep', 'prosent' eller 'heltall' per kolonnenavn."""
    kolonnetyper = kolonnetyper or {}
    status_kolonner = status_kolonner or []
    ws = wb.create_sheet(_sikkert_arknavn(arknavn))

    start_rad = 1
    if undertittel:
        ws.cell(row=1, column=1, value=undertittel).font = UNDERTITTEL_FONT
        start_rad = 3

    header_rad = start_rad
    for j, kolonne in enumerate(df.columns, start=1):
        celle = ws.cell(row=header_rad, column=j, value=str(kolonne))
        celle.font = HEADER_FONT
        celle.fill = HEADER_FILL
        celle.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, (_, rad) in enumerate(df.iterrows(), start=header_rad + 1):
        for j, kolonne in enumerate(df.columns, start=1):
            verdi = rad[kolonne]
            if isinstance(verdi, float) and pd.isna(verdi):
                verdi = None
            if verdi is pd.NaT:
                verdi = None
            celle = ws.cell(row=i, column=j, value=verdi)
            celle.font = BODY_FONT
            celle.alignment = Alignment(vertical="top", wrap_text=False)
            kotype = kolonnetyper.get(kolonne)
            if kotype == "dato" and isinstance(verdi, (dt.date, dt.datetime, pd.Timestamp)):
                celle.number_format = "dd.mm.yyyy"
            elif kotype == "beloep":
                celle.number_format = BELOPFORMAT
            elif kotype == "prosent":
                celle.number_format = PROSENTFORMAT
            elif kotype == "heltall":
                celle.number_format = HELTALLFORMAT
            if kolonne in status_kolonner:
                farge = _statusfarge(verdi)
                if farge:
                    celle.fill = farge

    siste_rad = header_rad + len(df)
    siste_kol = len(df.columns)
    if siste_rad > header_rad and siste_kol > 0:
        tabellomraade = (
            f"{get_column_letter(1)}{header_rad}:{get_column_letter(siste_kol)}{siste_rad}"
        )
        tabellnavn = "Tbl_" + "".join(c for c in arknavn if c.isalnum())[:20] or "Tabell"
        try:
            tabell = Table(displayName=tabellnavn, ref=tabellomraade)
            tabell.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False
            )
            ws.add_table(tabell)
        except ValueError:
            ws.auto_filter.ref = tabellomraade

    ws.freeze_panes = ws.cell(row=header_rad + 1, column=1).coordinate

    for j, kolonne in enumerate(df.columns, start=1):
        lengde = max([len(str(kolonne))] + [len(str(v)) for v in df[kolonne].astype(str).head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(lengde + 2, 10), 60)


def skriv_lederoppsummering(wb: Workbook, arknavn: str, linjer: list[tuple[str, str]]) -> None:
    """Fritekst-oppsummering (overskrift + verdi per linje), ingen sammenslåtte
    celler - hver linje er én rad med to kolonner (Punkt / Verdi)."""
    ws = wb.create_sheet(_sikkert_arknavn(arknavn))
    ws.cell(row=1, column=1, value="Punkt").font = HEADER_FONT
    ws.cell(row=1, column=1).fill = HEADER_FILL
    ws.cell(row=1, column=2, value="Verdi/kommentar").font = HEADER_FONT
    ws.cell(row=1, column=2).fill = HEADER_FILL
    for i, (punkt, verdi) in enumerate(linjer, start=2):
        c1 = ws.cell(row=i, column=1, value=punkt)
        c1.font = Font(name=FONT, size=10, bold=punkt.isupper())
        c2 = ws.cell(row=i, column=2, value=verdi)
        c2.font = BODY_FONT
        c2.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 90
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:B{len(linjer) + 1}"


def lag_arbeidsbok() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def lagre(wb: Workbook, sti: Path) -> None:
    sti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(sti)
