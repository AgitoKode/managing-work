"""Innlesing av RecMan-arbeidsboken uten å endre originalfilen.

Går gjennom alle synlige ark, finner automatisk hvilken rad som inneholder
kolonneoverskriftene (ikke nødvendigvis rad 1), og identifiserer hvilke ark
som faktisk inneholder arbeidsdata (kontra f.eks. et "Søkekriterier"-ark med
eksportmetadata).
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd

from aga_lib.column_mapping import FELT_ALIASER, bygg_kolonnemapping

MAKS_RADER_FOR_HEADERSOEK = 10
MIN_STANDARDFELT_FOR_DATAARK = 5


@dataclass
class ArkResultat:
    arknavn: str
    header_rad_index: int  # 0-basert
    df: pd.DataFrame
    kolonnemapping: "object"
    er_dataark: bool
    kolonner_med_mellomrom: list[str]
    kolonner_lagret_som_tekst_men_ser_numerisk_ut: list[str]


def _finn_header_rad(rader: list[list]) -> int:
    """Returnerer 0-basert indeks for raden som mest sannsynlig er
    kolonneoverskrifter, basert på hvor mange celler som matcher kjente
    feltalias etter normalisering."""
    kjente_aliaser = {a for aliaser in FELT_ALIASER.values() for a in aliaser}

    def score(rad: list) -> int:
        treff = 0
        for celle in rad:
            if celle is None:
                continue
            norm = str(celle).strip().lower().replace(".", " ").replace("-", " ")
            norm = " ".join(norm.split())
            if norm in kjente_aliaser:
                treff += 1
        return treff

    beste_index, beste_score = 0, -1
    for i, rad in enumerate(rader[:MAKS_RADER_FOR_HEADERSOEK]):
        s = score(rad)
        if s > beste_score:
            beste_index, beste_score = i, s
    return beste_index


def les_arbeidsbok(sti: Path) -> dict[str, ArkResultat]:
    """Leser alle synlige ark. Åpner arbeidsboken read-only og skriver aldri
    tilbake til filen - originalfilen er dermed garantert uendret."""
    wb = openpyxl.load_workbook(sti, data_only=True, read_only=True)
    resultater: dict[str, ArkResultat] = {}
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            rader_forhaandsvisning = []
            for i, rad in enumerate(ws.iter_rows(min_row=1, max_row=MAKS_RADER_FOR_HEADERSOEK, values_only=True)):
                rader_forhaandsvisning.append(list(rad))
                if i >= MAKS_RADER_FOR_HEADERSOEK - 1:
                    break
            if not rader_forhaandsvisning:
                continue
            header_index = _finn_header_rad(rader_forhaandsvisning)
            header = [c for c in rader_forhaandsvisning[header_index]]

            kolonner_med_mellomrom = [
                str(c) for c in header if c is not None and str(c) != str(c).strip()
            ]

            data_rader = []
            for i, rad in enumerate(ws.iter_rows(min_row=header_index + 2, values_only=True)):
                data_rader.append(list(rad))

            bredde = len(header)
            data_rader = [rad[:bredde] + [None] * (bredde - len(rad)) for rad in data_rader]
            df = pd.DataFrame(data_rader, columns=[str(h) if h is not None else f"_ukjent_{i}" for i, h in enumerate(header)])

            mapping = bygg_kolonnemapping(list(df.columns))
            er_dataark = len(mapping.standard_til_original) >= MIN_STANDARDFELT_FOR_DATAARK

            tekst_men_numerisk = []
            for kolonne in df.columns:
                serie = df[kolonne].dropna()
                if serie.empty:
                    continue
                if serie.map(lambda v: isinstance(v, str)).all():
                    proev = serie.head(50).str.replace(",", ".", regex=False).str.replace(" ", "", regex=False)
                    if proev.map(lambda v: v.replace(".", "", 1).replace("-", "", 1).isdigit() if v else False).mean() > 0.8:
                        tekst_men_numerisk.append(str(kolonne))

            resultater[ws.title] = ArkResultat(
                arknavn=ws.title,
                header_rad_index=header_index,
                df=df,
                kolonnemapping=mapping,
                er_dataark=er_dataark,
                kolonner_med_mellomrom=kolonner_med_mellomrom,
                kolonner_lagret_som_tekst_men_ser_numerisk_ut=tekst_men_numerisk,
            )
    finally:
        wb.close()
    return resultater
